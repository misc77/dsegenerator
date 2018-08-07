from docx import Document
import xml.etree.ElementTree as ET
from checklist import Checklist

def parseTable(xmlElement, checklistDocument):
    pos = int(xmlElement.attrib.get("tabelle"))
    table = checklistDocument.tables[pos-1]
    tableObject = {}
    
    for child in xmlElement:
        zeile = int(child.attrib.get("zeile"))-1
        spalte = int(child.attrib.get("spalte"))-1
        print (str(zeile) + " " + str(spalte))

        if child.attrib.get("wordType") == "table":
            tableObject[child.tag] = parseTable(child, table.cell(zeile,spalte))
        else:
            tableObject[child.tag] = table.cell(zeile,spalte).text
            print (child.tag, child.attrib) 
            print (table.cell(zeile,spalte).text)
    
    return tableObject

def parseChecklist(xmlFile, checklistFile):
    tree = ET.parse(xmlFile)
    root = tree.getroot()
    checklistDocument = Document(checklistFile)
    checklistObject = Checklist()
    checklistObject.xmlVersion = root.attrib.get("version")

    for elem in root:
        print(elem.attrib.get("wordType"))
        if elem.attrib.get("wordType") == "table":
            checklistObject.addCategory(elem.tag, parseTable(elem, checklistDocument))

    """ setting instance variables """
    checklistObject.setInstanceVariables()
    print("checklistobject...")
    print(checklistObject.getCategoryList()) 
    print(checklistObject.wordVersion + " " + checklistObject.xmlVersion)


parseChecklist(     "C:/Users/misc2/Documents/PyProjects/DSEGenerator/test/data/checkliste.xml"
                ,   "C:/Users/misc2/Documents/PyProjects/DSEGenerator/test/data/checkliste.docx" )
